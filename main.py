import os
import io
import openpyxl
import jwt
from typing import List, Optional
from datetime import datetime, timedelta, timezone

from fastapi import FastAPI, Depends, HTTPException, status, Request, Response
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.security import OAuth2PasswordRequestForm
from sqlalchemy.orm import Session
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
from dotenv import load_dotenv

import models, schemas, auth
from database import engine, get_db, SessionLocal

load_dotenv()
ENVIRONMENT = os.getenv("ENVIRONMENT", "development")

models.Base.metadata.create_all(bind=engine)

app = FastAPI(
    title="Controle de Divergências",
    docs_url=None if ENVIRONMENT == "production" else "/docs",
    redoc_url=None if ENVIRONMENT == "production" else "/redoc"
)

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

@app.middleware("http")
async def add_security_headers(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    if ENVIRONMENT == "production":
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response

def get_current_user(request: Request, db: Session = Depends(get_db)):
    token = request.cookies.get("access_token")
    credentials_exception = HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Credenciais inválidas ou sessão expirada",
    )
    if not token:
        raise credentials_exception
        
    token = token.replace("Bearer ", "")
    try:
        payload = jwt.decode(token, auth.SECRET_KEY, algorithms=[auth.ALGORITHM])
        username: str = payload.get("sub")
        if username is None:
            raise credentials_exception
    except jwt.PyJWTError:
        raise credentials_exception
        
    user = db.query(models.Usuario).filter(models.Usuario.username == username).first()
    if user is None:
        raise credentials_exception
    return user

@app.on_event("startup")
def startup_event():
    db = SessionLocal()
    first_user = os.getenv("FIRST_SUPERUSER_USERNAME")
    first_pass = os.getenv("FIRST_SUPERUSER_PASSWORD")
    
    if first_user and first_pass:
        usuario_existe = db.query(models.Usuario).filter(models.Usuario.username == first_user).first()
        if not usuario_existe:
            senha_criptografada = auth.get_password_hash(first_pass)
            novo_usuario = models.Usuario(username=first_user, hashed_password=senha_criptografada)
            db.add(novo_usuario)
            db.commit()
    db.close()

@app.post("/api/login")
@limiter.limit("5/minute")
def login(request: Request, response: Response, form_data: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(models.Usuario).filter(models.Usuario.username == form_data.username).first()
    if not user or not auth.verify_password(form_data.password, user.hashed_password):
        raise HTTPException(status_code=400, detail="Usuário ou senha incorretos")
    
    access_token = auth.create_access_token(data={"sub": user.username})
    
    response.set_cookie(
        key="access_token",
        value=f"Bearer {access_token}",
        httponly=True,
        secure=(ENVIRONMENT == "production"),
        samesite="strict",
        max_age=auth.ACCESS_TOKEN_EXPIRE_MINUTES * 60
    )
    return {"message": "Login efetuado com sucesso"}

@app.post("/api/logout")
def logout(response: Response):
    response.delete_cookie("access_token")
    return {"message": "Sessão terminada"}

@app.get("/api/estatisticas/")
def obter_estatisticas(request: Request, db: Session = Depends(get_db), current_user: models.Usuario = Depends(get_current_user)):
    total = db.query(models.Divergencia).count()
    faltas = db.query(models.Divergencia).filter(models.Divergencia.categoria == 'Falta').count()
    sobras = db.query(models.Divergencia).filter(models.Divergencia.categoria == 'Sobra').count()
    defeitos = db.query(models.Divergencia).filter(models.Divergencia.categoria == 'Defeito').count()
    return {"total": total, "faltas": faltas, "sobras": sobras, "defeitos": defeitos}

@app.post("/api/divergencias/", response_model=schemas.DivergenciaResponse)
@limiter.limit("30/minute")
def criar_divergencia(request: Request, div: schemas.DivergenciaCreate, db: Session = Depends(get_db), current_user: models.Usuario = Depends(get_current_user)):
    db_div = models.Divergencia(**div.model_dump())
    db.add(db_div)
    db.commit()
    db.refresh(db_div)
    return db_div

@app.get("/api/divergencias/", response_model=List[schemas.DivergenciaResponse])
def listar_divergencias(request: Request, skip: int = 0, limit: int = 50, sku: Optional[str] = None, categoria: Optional[str] = None, db: Session = Depends(get_db), current_user: models.Usuario = Depends(get_current_user)):
    if limit > 100: limit = 100
    
    query = db.query(models.Divergencia)
    if sku:
        query = query.filter(models.Divergencia.sku.icontains(sku))
    if categoria:
        query = query.filter(models.Divergencia.categoria == categoria)
    return query.offset(skip).limit(limit).all()

@app.put("/api/divergencias/{div_id}", response_model=schemas.DivergenciaResponse)
def atualizar_divergencia(request: Request, div_id: int, div: schemas.DivergenciaCreate, db: Session = Depends(get_db), current_user: models.Usuario = Depends(get_current_user)):
    db_div = db.query(models.Divergencia).filter(models.Divergencia.id == div_id).first()
    if not db_div:
        raise HTTPException(status_code=404, detail="Registro não encontrado")
    for key, value in div.model_dump().items():
        setattr(db_div, key, value)
    db.commit()
    db.refresh(db_div)
    return db_div

@app.delete("/api/divergencias/{div_id}")
def deletar_divergencia(request: Request, div_id: int, db: Session = Depends(get_db), current_user: models.Usuario = Depends(get_current_user)):
    db_div = db.query(models.Divergencia).filter(models.Divergencia.id == div_id).first()
    if not db_div:
        raise HTTPException(status_code=404, detail="Registro não encontrado")
    db.delete(db_div)
    db.commit()
    return {"message": "Registro excluído"}

@app.get("/api/exportar/excel")
@limiter.limit("2/minute")
def exportar_excel(request: Request, categoria: Optional[str] = None, db: Session = Depends(get_db), current_user: models.Usuario = Depends(get_current_user)):
    fuso_br = timezone(timedelta(hours=-3))
    agora = datetime.now(fuso_br)
    data_hora_br = agora.strftime("%d/%m/%Y às %H:%M:%S")
    data_hora_arquivo = agora.strftime("%d-%m-%Y_%H-%M")

    query = db.query(models.Divergencia)
    if categoria:
        query = query.filter(models.Divergencia.categoria == categoria)
        
    divergencias = query.limit(5000).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Divergências"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    
    titulo_relatorio = "RELATÓRIO DE DIVERGÊNCIAS DE PEÇAS"
    if categoria:
        titulo_relatorio += f" - {categoria.upper()}"

    ws.merge_cells('A1:F1')
    title_cell = ws['A1']
    title_cell.value = titulo_relatorio
    title_cell.font = Font(size=14, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
    ws.row_dimensions[1].height = 30

    headers = ["ID", "SKU (Part Number)", "Quantidade", "Categoria", "Subcategoria", "Valor Compra"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="D50000", end_color="D50000", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    borda_estilo = Side(style='thin', color="CCCCCC")
    borda = Border(left=borda_estilo, right=borda_estilo, top=borda_estilo, bottom=borda_estilo)

    for row_num, d in enumerate(divergencias, 3):
        ws.cell(row=row_num, column=1, value=d.id).border = borda
        ws.cell(row=row_num, column=2, value=d.sku).border = borda
        ws.cell(row=row_num, column=3, value=d.quantidade).border = borda
        ws.cell(row=row_num, column=4, value=d.categoria).border = borda
        ws.cell(row=row_num, column=5, value=d.subcategoria or '-').border = borda
        
        valor_str = f"R$ {d.valor_compra:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if d.valor_compra else "-"
        ws.cell(row=row_num, column=6, value=valor_str).border = borda
        
        ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=4).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=6).alignment = Alignment(horizontal="center", vertical="center")

    ultima_linha = ws.max_row + 1 if divergencias else 3
    
    ws.merge_cells(start_row=ultima_linha, start_column=1, end_row=ultima_linha, end_column=6)
    rodape_cell = ws.cell(row=ultima_linha, column=1)
    rodape_cell.value = f"Documento gerado em: {data_hora_br}"
    rodape_cell.font = Font(italic=True, size=10, color="555555")
    rodape_cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[ultima_linha].height = 25

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 20

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    nome_filtro = f"_{categoria.lower()}" if categoria else ""
    nome_arquivo = f"relatorio_divergencias{nome_filtro}_{data_hora_arquivo}.xlsx"
    
    return StreamingResponse(
        iter([output.getvalue()]), 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
        headers={"Content-Disposition": f"attachment; filename={nome_arquivo}"}
    )

app.mount("/", StaticFiles(directory="static", html=True), name="static")